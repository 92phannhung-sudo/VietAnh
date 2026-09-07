import unittest
from docx import Document
from docx.shared import Mm, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from scripts.format_m7_document import apply_page_setup, apply_typography

class TestM7PageSetupAndTypography(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.p = self.doc.add_paragraph("Đoạn văn thử nghiệm nội dung thân bài thuyết minh sáng kiến.")
    
    def test_page_margins(self):
        apply_page_setup(self.doc)
        section = self.doc.sections[0]
        self.assertAlmostEqual(section.top_margin.mm, 20.0, places=1)
        self.assertAlmostEqual(section.bottom_margin.mm, 20.0, places=1)
        self.assertAlmostEqual(section.left_margin.mm, 30.0, places=1)
        self.assertAlmostEqual(section.right_margin.mm, 15.0, places=1)

    def test_typography(self):
        apply_typography(self.doc)
        p = self.doc.paragraphs[0]
        self.assertEqual(p.runs[0].font.name, "Times New Roman")
        self.assertEqual(p.runs[0].font.size, Pt(13))
        self.assertAlmostEqual(p.paragraph_format.first_line_indent.cm, 1.27, places=2)

    def test_clean_table_data_and_format_tables(self):
        table = self.doc.add_table(rows=3, cols=5)
        # Hàng tiêu đề
        table.cell(0, 0).text = "STT"
        table.cell(0, 1).text = "Hạng mục"
        table.cell(0, 3).text = "Đơn giá"
        table.cell(0, 4).text = "Thành tiền"
        # Hàng dữ liệu 1 có typo ".."
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "Webcam Logitech 920e"
        table.cell(1, 3).text = "1.800.000"
        table.cell(1, 4).text = "1.800..000"
        # Hàng dữ liệu 2 có typo ".."
        table.cell(2, 0).text = "2"
        table.cell(2, 1).text = "Bàn đạp chân"
        table.cell(2, 3).text = "290..000"
        table.cell(2, 4).text = "290..000"

        from scripts.format_m7_document import clean_table_data, format_tables
        clean_table_data(self.doc)
        format_tables(self.doc)

        self.assertEqual(table.cell(1, 4).text, "1.800.000")
        self.assertEqual(table.cell(2, 3).text, "290.000")
        self.assertEqual(table.cell(2, 4).text, "290.000")
        self.assertEqual(table.cell(0, 0).paragraphs[0].runs[0].font.bold, True)

    def test_insert_images_and_captions(self):
        doc = Document()
        p = doc.add_paragraph("Ảnh 1 : Mã định danh")
        from scripts.format_m7_document import insert_images_and_captions
        images_dir = "/Volumes/DATA/NguyenVietAnh/docs/report"
        insert_images_and_captions(doc, images_dir)
        xml_str = doc._body._element.xml
        self.assertIn("pic:pic", xml_str)
        # Kiểm tra caption chuẩn đã được tạo
        captions = [para.text for para in doc.paragraphs if "Ảnh 1:" in para.text]
        self.assertTrue(len(captions) > 0)

    def test_split_soft_break_paragraphs(self):
        doc = Document()
        p = doc.add_paragraph("Bước 1. Tiếp nhận và đối chiếu bệnh phẩm\nNgười nhận tiếp nhận bệnh phẩm từ người giao; đối chiếu thông tin giữa LIS.")
        from scripts.format_m7_document import split_soft_break_paragraphs, apply_typography
        split_soft_break_paragraphs(doc)
        apply_typography(doc)

        self.assertEqual(len(doc.paragraphs), 2)
        p1 = doc.paragraphs[0]
        p2 = doc.paragraphs[1]

        # Paragraph 1: Tiêu đề bước
        self.assertEqual(p1.text, "Bước 1. Tiếp nhận và đối chiếu bệnh phẩm")
        self.assertEqual(p1.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertTrue(p1.runs[0].font.bold)
        self.assertTrue(p1.paragraph_format.keep_with_next)
        self.assertNotIn("\n", p1.text)

        # Paragraph 2: Thân nội dung bước
        self.assertEqual(p2.text, "Người nhận tiếp nhận bệnh phẩm từ người giao; đối chiếu thông tin giữa LIS.")
        self.assertEqual(p2.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertFalse(p2.runs[0].font.bold)
        self.assertNotIn("\n", p2.text)

if __name__ == '__main__':
    unittest.main()
