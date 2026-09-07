import unittest
import os
import re
from docx import Document

class TestDossierConsistency(unittest.TestCase):
    def setUp(self):
        self.base_dir = "/Volumes/DATA/NguyenVietAnh/docs/report"
        self.m1_path = os.path.join(self.base_dir, "Bản sao M1 Phieu dang ky.docx")
        self.m8_path = os.path.join(self.base_dir, "Bản sao M8 Xac nhan Ty le dong gop Cong trinh (1).docx")
        self.m7_path = os.path.join(self.base_dir, "M7_Thuyet_minh_Cong_trinh_ChuanHoa.docx")

    def test_m1_date_and_author(self):
        self.assertTrue(os.path.exists(self.m1_path), "File M1 không tồn tại")
        doc = Document(self.m1_path)
        full_text = " ".join(p.text for p in doc.paragraphs)
        
        # Kiểm tra chính tả tên
        self.assertNotIn("Nguyễn Viêt Anh", full_text, "Tên tác giả trong M1 bị lỗi chính tả 'Nguyễn Viêt Anh'")
        self.assertIn("Nguyễn Việt Anh", full_text)
        
        # Kiểm tra mốc năm ký
        self.assertIn("2026", full_text, "M1 phải có mốc năm ký 2026")
        self.assertNotIn("2025", full_text, "M1 không được lẫn mốc năm 2025")

    def test_m8_date_consistency(self):
        self.assertTrue(os.path.exists(self.m8_path), "File M8 không tồn tại")
        doc = Document(self.m8_path)
        full_text = " ".join(p.text for p in doc.paragraphs)
        
        # M8 phải có mốc năm ký 2026, không được là 2025
        self.assertIn("2026", full_text, "M8 phải thống nhất mốc năm 2026")
        self.assertNotIn("2025", full_text, "M8 không được chứa mốc năm cũ 2025")

    def test_m7_date_and_author_table(self):
        self.assertTrue(os.path.exists(self.m7_path), "File M7 chuẩn hóa không tồn tại")
        doc = Document(self.m7_path)
        
        # Kiểm tra ngày tháng cuối bài
        date_paras = [p.text.strip() for p in doc.paragraphs if "ngày 25 tháng 8" in p.text]
        self.assertTrue(len(date_paras) > 0, "Không tìm thấy dòng ngày tháng trong M7")
        self.assertIn("2026", date_paras[0], "M7 phải có mốc ngày 25 tháng 8 năm 2026")
        self.assertNotIn("2025", date_paras[0], "M7 không được để mốc 2025")

        # Kiểm tra thông tin tác giả trong Table 0
        self.assertTrue(len(doc.tables) > 0, "M7 phải có bảng biểu")
        t0 = doc.tables[0]
        self.assertTrue(len(t0.rows) >= 3, "Bảng tác giả M7 phải có ít nhất 3 hàng (1 header + 2 tác giả)")
        
        # Hàng 2 là Nguyễn Phúc Đẳng
        row_dang_text = " ".join(c.text.strip() for c in t0.rows[2].cells)
        self.assertIn("Nguyễn Phúc Đẳng", row_dang_text)
        self.assertIn("0866085675", row_dang_text, "SĐT của Nguyễn Phúc Đẳng trong M7 phải là 0866085675")
        self.assertNotIn("0969577630", row_dang_text, "Không được copy SĐT của Nguyễn Việt Anh sang dòng Nguyễn Phúc Đẳng")
        self.assertIn("2002", row_dang_text, "Năm sinh của Nguyễn Phúc Đẳng phải là 2002")

if __name__ == '__main__':
    unittest.main()
